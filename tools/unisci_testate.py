"""Unisce i CSV regionali in data/testate.csv e data/testate.xlsx (per la revisione).

Uso: .venv/bin/python tools/unisci_testate.py
"""
import csv
from collections import Counter
from pathlib import Path
from urllib.parse import urlparse

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

BASE = Path(__file__).resolve().parent.parent
REGIONI = ["nordovest", "nordest", "centro", "sud", "isole"]
COLONNE = ["nome", "url_homepage", "feed_rss", "stato", "citta", "tipo", "gruppo",
           "lingua", "paywall", "homepage_status", "note"]


def chiave_url(url):
    """Dominio senza www più percorso: distingue le sezioni (es. latimes.com/espanol)."""
    u = urlparse(url.strip())
    host = u.netloc.lower()
    host = host[4:] if host.startswith("www.") else host
    return host + u.path.rstrip("/").lower()


def carica():
    righe, visti, doppioni, feed_visti = [], set(), [], set()
    for regione in REGIONI:
        path = BASE / "data" / "regioni" / f"{regione}.csv"
        if not path.exists():
            print(f"manca {path.name}")
            continue
        with open(path, newline="", encoding="utf-8") as f:
            for r in csv.DictReader(f):
                r = {k: (r.get(k) or "").strip() for k in COLONNE}
                chiave = chiave_url(r["url_homepage"])
                if chiave in visti:
                    doppioni.append(r["nome"])
                    continue
                visti.add(chiave)
                feed = r["feed_rss"].strip()
                if feed and feed in feed_visti:
                    # stesso feed usato per più edizioni: si tiene la testata senza feed (userà la homepage)
                    r["note"] = (r["note"] + "; feed condiviso con un'altra edizione").lstrip("; ")
                    r["feed_rss"] = ""
                elif feed:
                    feed_visti.add(feed)
                r["regione"] = regione
                r["includi"] = "si"
                righe.append(r)
    righe.sort(key=lambda r: (r["stato"], r["tipo"], r["nome"].lower()))
    return righe, doppioni


VERIFICA = ["metodo", "status_usa", "feed_ok", "elementi", "sitemap_news"]


def aggiungi_verifica(righe):
    """Se esiste data/verifica_usa.csv, aggiunge l'esito della verifica da USA."""
    path = BASE / "data" / "verifica_usa.csv"
    if not path.exists():
        return []
    with open(path, newline="", encoding="utf-8") as f:
        esiti = {chiave_url(e["url_homepage"]): e for e in csv.DictReader(f)}
    for r in righe:
        e = esiti.get(chiave_url(r["url_homepage"]), {})
        for c in VERIFICA:
            r[c] = e.get(c, "")
        if r["metodo"] == "nessuno":  # bloccate anche da USA: si ripiega su Google News
            r["metodo"] = "google_news"
    return VERIFICA


def scrivi(righe, extra):
    campi = ["includi", "regione"] + COLONNE + extra
    with open(BASE / "data" / "testate.csv", "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=campi)
        w.writeheader()
        w.writerows(righe)

    wb = Workbook()
    ws = wb.active
    ws.title = "Testate"
    ws.append(campi)
    for r in righe:
        ws.append([r[c] for c in campi])
    for cella in ws[1]:
        cella.font = Font(bold=True, color="FFFFFF")
        cella.fill = PatternFill("solid", fgColor="1F3A5F")
    larghezze = {"includi": 8, "regione": 10, "nome": 34, "url_homepage": 34, "feed_rss": 40,
                 "stato": 6, "citta": 18, "tipo": 18, "gruppo": 18, "lingua": 7,
                 "paywall": 8, "homepage_status": 8, "note": 60, "metodo": 14,
                 "status_usa": 9, "feed_ok": 40, "elementi": 9, "sitemap_news": 40}
    for i, c in enumerate(campi, 1):
        ws.column_dimensions[get_column_letter(i)].width = larghezze[c]
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = ws.dimensions

    riepilogo = wb.create_sheet("Riepilogo")
    for titolo, chiave in [("Metodo di raccolta (da USA)", "metodo"), ("Per stato", "stato"), ("Per tipo", "tipo"),
                           ("Per gruppo", "gruppo"), ("Per lingua", "lingua")]:
        riepilogo.append([titolo])
        riepilogo[riepilogo.max_row][0].font = Font(bold=True)
        for valore, n in Counter(r.get(chiave, "") for r in righe).most_common():
            riepilogo.append([valore, n])
        riepilogo.append([])
    riepilogo.column_dimensions["A"].width = 28
    wb.save(BASE / "data" / "testate.xlsx")


if __name__ == "__main__":
    righe, doppioni = carica()
    extra = aggiungi_verifica(righe)
    scrivi(righe, extra)
    print(f"{len(righe)} testate; doppioni rimossi: {doppioni or 'nessuno'}")
    print("stati coperti:", len({r['stato'] for r in righe}))
    print(Counter(r["tipo"] for r in righe).most_common())
    print(Counter(r["lingua"] for r in righe).most_common())
